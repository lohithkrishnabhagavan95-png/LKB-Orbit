"""Wires everything together: tray icon, the drag+Shift watcher, the radial menu,
background conversion jobs, and desktop notifications. Single-instance: a second
launch hands its files to the running copy and exits."""
from __future__ import annotations

import json
import os
import sys
import traceback
from pathlib import Path

from PySide6.QtCore import QObject, QRunnable, Qt, QThreadPool, QTimer, Signal
from PySide6.QtGui import QAction, QColor, QCursor, QFont, QIcon, QPainter, QPixmap
from PySide6.QtNetwork import QLocalServer, QLocalSocket
from PySide6.QtWidgets import QApplication, QDialog, QMenu, QSystemTrayIcon, QFileDialog

from . import engine, integration
from .dialogs import PagePicker
from .menu import RadialMenu
from .trigger import DragWatcher
from .util import APP_NAME, CONFIG, UserError, file_size, human_size, resolve_shortcuts


def ipc_name() -> str:
    return f"{APP_NAME}-" + (os.environ.get("USERNAME") or os.environ.get("USER") or "user")


def resource_path(rel: str) -> str:
    """Locate a bundled resource both in dev and inside a PyInstaller onefile exe."""
    base = getattr(sys, "_MEIPASS", os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    return os.path.join(base, rel)


_ICON: QIcon | None = None


def make_icon() -> QIcon:
    global _ICON
    if _ICON is None:
        path = resource_path(os.path.join("quickconvert", "assets", "icon.png"))
        if os.path.exists(path):
            _ICON = QIcon(path)
        else:  # fallback: draw a simple orbit mark
            pm = QPixmap(64, 64)
            pm.fill(Qt.transparent)
            p = QPainter(pm)
            p.setRenderHint(QPainter.Antialiasing)
            p.setBrush(QColor(37, 99, 235)); p.setPen(Qt.NoPen)
            p.drawEllipse(2, 2, 60, 60)
            p.setPen(QColor(255, 255, 255)); p.setFont(QFont("Segoe UI", 26, QFont.Bold))
            p.drawText(pm.rect(), Qt.AlignCenter, "O")
            p.end()
            _ICON = QIcon(pm)
    return _ICON


# ---------------------------------------------------------------- worker
class _Sig(QObject):
    done = Signal(list, float)  # outputs, seconds
    failed = Signal(str)
    progress = Signal(int, int)  # done, total


class Job(QRunnable):
    def __init__(self, cat, ring, label, paths, params=None):
        super().__init__()
        self.cat, self.ring, self.label, self.paths = cat, ring, label, paths
        self.params = params or {}
        self.sig = _Sig()

    def run(self):
        import time
        t0 = time.time()
        try:
            outs = engine.run(self.cat, self.ring, self.label, self.paths, self.params,
                              on_progress=lambda d, t: self.sig.progress.emit(d, t))
            self.sig.done.emit([str(o) for o in outs], time.time() - t0)
        except UserError as e:
            self.sig.failed.emit(str(e))
        except Exception:
            traceback.print_exc()
            self.sig.failed.emit("Something went wrong with that conversion.")


class QuickConvert(QObject):
    def __init__(self, app: QApplication):
        super().__init__()
        self.app = app
        self.pool = QThreadPool.globalInstance()
        self._jobs: set[Job] = set()
        self.last_outputs: list[str] = []
        self._picking = False

        self.menu = RadialMenu()
        self.menu.action_chosen.connect(self._start)

        self.watcher = DragWatcher(self.menu.isVisible, lambda: self.menu.got_drag,
                                   lambda: self.menu.click_mode, lambda: self.menu.converting)
        self.watcher.set_key(CONFIG.get("trigger", "shift"))
        self.watcher.show_menu.connect(lambda pos: self.menu.popup(pos))
        self.watcher.toggle_ring.connect(self.menu.toggle)
        self.watcher.hide_menu.connect(self.menu.hide)
        self.watcher.start()

        self.tray = QSystemTrayIcon(make_icon(), self.app)
        self.tray.setToolTip(f"{APP_NAME} — drag a file, hold Shift")
        self.tray.setContextMenu(self._tray_menu())
        self.tray.activated.connect(self._tray_activated)
        self.tray.show()

    # ------------------------------------------------------------ tray
    def _tray_menu(self) -> QMenu:
        m = QMenu()
        m.addAction(f"{APP_NAME} — drag a file + Shift").setEnabled(False)
        m.addSeparator()
        m.addAction("Open a file…", self.pick_file)
        out = m.addMenu("Save output")
        for key, text in (("next", "Next to the original"), ("folder", f"In a {APP_NAME} subfolder")):
            a = QAction(text, out, checkable=True)
            a.setChecked(CONFIG["output"] == key)
            a.triggered.connect(lambda _c, k=key: CONFIG.__setitem__("output", k))
            out.addAction(a)
        trig = m.addMenu("Trigger key")
        for key in ("shift", "ctrl", "alt"):
            a = QAction(key.capitalize(), trig, checkable=True)
            a.setChecked(CONFIG["trigger"] == key)
            a.triggered.connect(lambda _c, k=key: self._set_trigger(k))
            trig.addAction(a)

        if integration.available():
            m.addSeparator()
            a1 = QAction("Start with Windows", m, checkable=True)
            a1.setChecked(integration.is_startup())
            a1.triggered.connect(self._toggle_startup)
            m.addAction(a1)
            a2 = QAction("Add to right-click menu", m, checkable=True)
            a2.setChecked(integration.has_context())
            a2.triggered.connect(self._toggle_context)
            m.addAction(a2)

        m.addSeparator()
        m.addAction("Quit", self.app.quit)
        return m

    def _toggle_startup(self, on: bool):
        try:
            integration.set_startup(on)
            self.tray.showMessage(APP_NAME, "Will start with Windows." if on else "Won't start with Windows.",
                                  make_icon(), 3000)
        except Exception as e:
            self.tray.showMessage(APP_NAME, f"Couldn't change startup: {e}", QSystemTrayIcon.Warning, 4000)

    def _toggle_context(self, on: bool):
        try:
            integration.set_context(on)
            self.tray.showMessage(APP_NAME, "Right-click entry added." if on else "Right-click entry removed.",
                                  make_icon(), 3000)
        except Exception as e:
            self.tray.showMessage(APP_NAME, f"Couldn't update right-click menu: {e}", QSystemTrayIcon.Warning, 4000)

    def _set_trigger(self, key: str):
        CONFIG["trigger"] = key
        self.watcher.set_key(key)

    def _tray_activated(self, reason):
        if reason in (QSystemTrayIcon.Trigger, QSystemTrayIcon.DoubleClick):
            # Defer out of the signal handler, and the guard in pick_file stops a
            # double-click (Trigger + DoubleClick) opening two native dialogs at once.
            QTimer.singleShot(0, self.pick_file)

    def pick_file(self):
        if self._picking:
            return
        self._picking = True
        try:
            files, _ = QFileDialog.getOpenFileNames(None, f"{APP_NAME} — choose a file")
        except Exception:
            files = []
        finally:
            self._picking = False
        if files:
            self.open_files(files)

    # ------------------------------------------------------------ run
    def open_files(self, files: list[str]):
        files = resolve_shortcuts([f for f in files if os.path.isfile(f)])
        if files:
            self.menu.popup(QCursor.pos(), paths=files)

    def _start(self, cat, ring, label, paths):
        params = {}
        if ring == "tools" and label in engine.NEEDS_PAGE_PICK:
            params = self._ask_pages(cat, label, paths)
            if params is None:   # cancelled
                self.menu.hide()
                return
        job = Job(cat, ring, label, paths, params)
        job.setAutoDelete(False)
        self._jobs.add(job)
        job.sig.done.connect(lambda outs, secs, j=job: self._finished(outs, secs, j))
        job.sig.failed.connect(lambda msg, j=job: self._error(msg, j))
        job.sig.progress.connect(self.menu.set_progress)
        op = f"Converting to {label}…" if ring == "convert" else f"{label}…"
        self.menu.begin_convert(op)
        self.pool.start(job)
        self.tray.setToolTip(f"{APP_NAME} — converting…")

    def _ask_pages(self, cat, label, paths) -> dict | None:
        """Show the page/slide picker. Returns params dict, {} to use all, or None if cancelled."""
        unit = "slide" if cat == "slide" else "page"
        try:
            n = engine.page_count(cat, paths[0])
        except UserError as e:
            self.tray.showMessage(APP_NAME, str(e), QSystemTrayIcon.Warning, 4000)
            return None
        except Exception:
            self.tray.showMessage(APP_NAME, "Couldn't read that file.", QSystemTrayIcon.Warning, 4000)
            return None
        if n <= 1:
            return {}
        dlg = PagePicker(n, f"{label} — {os.path.basename(paths[0])}", unit)
        dlg.setWindowFlag(Qt.WindowStaysOnTopHint, True)
        if dlg.exec() != QDialog.Accepted:
            return None
        sel = dlg.selected()
        return {"pages": sel} if sel else None

    def _finished(self, outputs, _secs, job):
        self._jobs.discard(job)
        self.last_outputs = outputs
        self.tray.setToolTip(f"{APP_NAME} — drag a file, hold Shift")
        self.menu.convert_done(outputs)
        if not outputs:
            return
        total = sum(file_size(o) for o in outputs)
        name = os.path.basename(outputs[0]) + (f" +{len(outputs) - 1} more" if len(outputs) > 1 else "")
        self.tray.showMessage(f"{APP_NAME}: done", f"{name}  ·  {human_size(total)}",
                              make_icon(), 4000)

    def _error(self, msg, job):
        self._jobs.discard(job)
        self.tray.setToolTip(f"{APP_NAME} — drag a file, hold Shift")
        self.menu.convert_fail(msg)
        self.tray.showMessage(f"{APP_NAME}: couldn't convert", msg, QSystemTrayIcon.Warning, 5000)


# ---------------------------------------------------------------- single instance
def _send_to_running(files: list[str], show: bool) -> bool:
    sock = QLocalSocket()
    sock.connectToServer(ipc_name())
    if not sock.waitForConnected(300):
        return False
    sock.write(json.dumps({"files": files, "show": show}).encode())
    sock.flush()
    sock.waitForBytesWritten(300)
    sock.disconnectFromServer()
    return True


def _listen(qc: QuickConvert):
    server = QLocalServer()
    QLocalServer.removeServer(ipc_name())
    server.listen(ipc_name())

    def on_conn():
        c = server.nextPendingConnection()
        if c and c.waitForReadyRead(300):
            try:
                msg = json.loads(bytes(c.readAll().data()).decode())
            except Exception:
                msg = {}
            if msg.get("files"):
                qc.open_files(msg["files"])
        c.deleteLater()

    server.newConnection.connect(on_conn)
    qc._server = server  # keep alive


def _selftest() -> int:
    """Headless sanity check used by packaging/CI: exercises every bundled engine
    library (Pillow, PyMuPDF, openpyxl, python-pptx) so a packaged exe self-verifies."""
    import tempfile
    import traceback
    from PIL import Image
    d = Path(tempfile.mkdtemp())
    try:
        png = d / "a.png"
        Image.new("RGB", (40, 30), (1, 2, 3)).save(png)
        assert engine.image_to(png, "JPG")[0].exists()          # Pillow
        pdf = engine.images_to_pdf([png])[0]
        assert engine.pdf_to_txt(pdf)[0].exists()                # PyMuPDF
        import openpyxl
        xf = d / "s.xlsx"
        wb = openpyxl.Workbook()
        wb.active.append(["x", 1])
        wb.save(xf)
        assert engine.sheet_to_csv(xf)[0].exists()               # openpyxl
        assert engine.sheet_to_pdf(xf)[0].exists()               # openpyxl + fitz.Story
        from pptx import Presentation
        pf = d / "p.pptx"
        pr = Presentation()
        for title in ("Slide One", "Slide Two"):
            s = pr.slides.add_slide(pr.slide_layouts[5])
            s.shapes.title.text = title
        pr.save(pf)
        assert engine.slide_to_txt(pf)[0].exists()               # python-pptx text
        assert engine.slide_to_pdf(pf)[0].exists()               # self-contained slide render -> PDF
        assert len(engine.slide_to_images(pf, "PNG")) == 2       # one PNG per slide
        if engine.HEIF_OK:
            assert engine.image_to(png, "HEIC")[0].exists()      # pillow-heif
        if engine.VTRACER_OK:
            assert engine.image_to(png, "SVG")[0].exists()       # vtracer (raster -> SVG)
        from ebooklib import epub
        bk = epub.EpubBook()
        bk.set_title("t")
        ch = epub.EpubHtml(title="c", file_name="c.xhtml")
        ch.content = "<html><body><p>hi</p></body></html>"
        bk.add_item(ch)
        bk.spine = [ch]
        bk.add_item(epub.EpubNcx())
        bk.add_item(epub.EpubNav())
        ep = d / "b.epub"
        epub.write_epub(str(ep), bk)
        assert engine.epub_to_txt(ep)[0].exists()                # ebooklib
        return 0
    except Exception:
        err = traceback.format_exc()
        try:
            (Path(tempfile.gettempdir()) / "lkborbit_selftest.log").write_text(err, encoding="utf-8")
        except Exception:
            pass
        print(err)
        return 1


def _install_crash_logging():
    """Write any fatal fault / uncaught exception to %TEMP% for diagnosis, since a
    windowed exe has no console."""
    import faulthandler
    import tempfile
    import traceback
    try:
        faulthandler.enable(open(os.path.join(tempfile.gettempdir(), "lkborbit_fault.log"), "w"))
    except Exception:
        pass

    def hook(t, v, tb):
        try:
            with open(os.path.join(tempfile.gettempdir(), "lkborbit_crash.log"), "w", encoding="utf-8") as f:
                f.write("".join(traceback.format_exception(t, v, tb)))
        except Exception:
            pass
        sys.__excepthook__(t, v, tb)

    sys.excepthook = hook


def main() -> int:
    argv = sys.argv
    if "--selftest" in argv:
        return _selftest()
    _install_crash_logging()

    files = [os.path.abspath(a) for a in argv[1:] if not a.startswith("-") and os.path.exists(a)]
    # Hand off to an already-running instance, then exit.
    if _send_to_running(files, show=not files):
        return 0

    app = QApplication(argv)
    app.setApplicationName(APP_NAME)
    app.setQuitOnLastWindowClosed(False)  # live in the tray
    if not QSystemTrayIcon.isSystemTrayAvailable():
        print("No system tray available.", file=sys.stderr)

    qc = QuickConvert(app)
    _listen(qc)
    if files:
        QTimer.singleShot(0, lambda: qc.open_files(files))
    return app.exec()
