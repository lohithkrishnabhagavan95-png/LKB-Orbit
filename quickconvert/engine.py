"""All file processing. No UI here, so it is directly testable. Every op takes
source path(s) and returns the list of output paths, written next to the source.
Originals are never modified or overwritten."""
from __future__ import annotations

import html
import io
import os
import sys
from pathlib import Path

from PIL import Image, ImageOps

from .util import CONFIG, UserError, sibling, sibling_folder, unique_path

try:
    import pymupdf as fitz
except ImportError:  # pragma: no cover
    import fitz  # type: ignore

# ------------------------------------------------------------------ targets
IMAGE_TARGETS = ["PNG", "JPG", "WEBP", "TIFF", "BMP", "GIF", "ICO", "PDF"]
PDF_TARGETS = ["JPG", "PNG", "DOCX", "TXT"]
DOCX_TARGETS = ["PDF", "TXT"]
TEXT_TARGETS = ["PDF", "DOCX"]
SHEET_TARGETS = ["PDF", "CSV", "XLSX"]
SLIDE_TARGETS = ["PDF", "PNG", "TXT"]

PIL_FORMAT = {"PNG": "PNG", "JPG": "JPEG", "WEBP": "WEBP", "TIFF": "TIFF",
              "BMP": "BMP", "GIF": "GIF", "ICO": "ICO", "PDF": "PDF"}
NO_ALPHA = {"JPG", "PDF", "BMP"}
EXT_TARGET = {".jpg": "JPG", ".jpeg": "JPG", ".png": "PNG", ".webp": "WEBP",
              ".tif": "TIFF", ".tiff": "TIFF", ".bmp": "BMP", ".gif": "GIF", ".ico": "ICO"}


# ------------------------------------------------------------------ io helpers
def atomic_save(write, out) -> None:
    """Write to a hidden temp file then rename, so Explorer never sees a partial file."""
    import os
    if not isinstance(out, (str, Path)):
        write(out)  # in-memory buffer
        return
    out = Path(out)
    tmp = out.with_name(f".~{out.name}.part")
    try:
        write(tmp)
        os.replace(tmp, out)
    finally:
        if tmp.exists():
            try:
                tmp.unlink()
            except OSError:
                pass


def open_image(path) -> Image.Image:
    im = Image.open(path)
    im.seek(0)
    im = ImageOps.exif_transpose(im)  # honour camera rotation
    if im.mode in ("P", "LA", "PA"):
        im = im.convert("RGBA")
    elif im.mode not in ("RGB", "RGBA", "L"):
        im = im.convert("RGB")
    return im


def flatten(im: Image.Image, bg=(255, 255, 255)) -> Image.Image:
    if im.mode == "RGBA":
        base = Image.new("RGB", im.size, bg)
        base.paste(im, mask=im.split()[3])
        return base
    return im.convert("RGB") if im.mode != "RGB" else im


def save_image(im: Image.Image, out, target: str, quality: int | None = None) -> None:
    quality = quality or int(CONFIG.get("jpg_quality", 90))
    target = target.upper()
    fmt = PIL_FORMAT[target]
    if target in NO_ALPHA:
        im = flatten(im)
    kw: dict = {}
    if target in ("JPG", "WEBP"):
        kw["quality"] = quality
    if target == "JPG":
        kw.update(optimize=True, progressive=True)
    if target == "PNG":
        kw["optimize"] = True
    if target == "ICO":
        side = min(256, max(im.size))
        im = ImageOps.contain(im.convert("RGBA"), (side, side))
        kw["sizes"] = [(s, s) for s in (16, 32, 48, 64, 128, 256) if s <= side]
    if target == "PDF":
        kw["resolution"] = 150.0
    atomic_save(lambda dst: im.save(dst, fmt, **kw), out)


def open_pdf(path):
    doc = fitz.open(path)
    if doc.needs_pass:
        doc.close()
        raise UserError("This PDF is password-protected. Remove the password first.")
    return doc


# ------------------------------------------------------------------ images
def image_to(src, target: str) -> list[Path]:
    target = target.upper()
    out = sibling(src, ext="." + target.lower())
    save_image(open_image(src), out, target)
    return [out]


def images_to_pdf(srcs: list) -> list[Path]:
    pages = [flatten(open_image(s)) for s in srcs]
    out = sibling(srcs[0], "Combined", ".pdf")
    pages[0].save(out, "PDF", resolution=150.0, save_all=True, append_images=pages[1:])
    return [out]


def compress_image(src, quality: int = 72) -> list[Path]:
    im = open_image(src)
    tgt = EXT_TARGET.get(Path(src).suffix.lower(), "PNG")
    if tgt == "PNG":  # lossy PNG via palette quantise (keeps alpha)
        colors = max(16, min(256, int(quality * 2.56)))
        q = im.convert("RGBA").quantize(colors=colors, method=Image.Quantize.FASTOCTREE,
                                        dither=Image.Dither.FLOYDSTEINBERG)
        out = sibling(src, "Compressed", ".png")
        atomic_save(lambda dst: q.save(dst, "PNG", optimize=True), out)
    elif tgt in ("BMP", "TIFF", "GIF", "ICO"):  # formats with no useful lossy mode -> JPG
        out = sibling(src, "Compressed", ".jpg")
        save_image(im, out, "JPG", quality=quality)
    else:
        out = sibling(src, "Compressed", Path(src).suffix.lower())
        save_image(im, out, tgt, quality=quality)
    return [out]


def resize_image(src, long_side: int = 1280) -> list[Path]:
    im = open_image(src)
    w, h = im.size
    if max(w, h) > long_side:
        k = long_side / max(w, h)
        im = im.resize((max(1, round(w * k)), max(1, round(h * k))), Image.LANCZOS)
    out = sibling(src, f"{im.size[0]}x{im.size[1]}")
    save_image(im, out, EXT_TARGET.get(Path(src).suffix.lower(), "PNG"))
    return [out]


# ------------------------------------------------------------------ pdf
def _pdf_pages_to_images(src, dpi: int) -> list[Image.Image]:
    imgs = []
    with open_pdf(src) as doc:
        for page in doc:
            pix = page.get_pixmap(dpi=dpi, alpha=False)
            imgs.append(Image.frombytes("RGB", (pix.width, pix.height), pix.samples))
    return imgs


def _save_pages(images: list[Image.Image], src, target: str, unit: str = "page") -> list[Path]:
    """Save one image per page/slide. If there's more than one, drop them into a
    folder named after the source so the original folder stays clean."""
    target = target.upper()
    ext = "." + target.lower()
    n = len(images)
    folder = sibling_folder(src) if n > 1 else None
    outs = []
    for i, im in enumerate(images):
        if folder is not None:
            out = unique_path(folder / f"{unit} {i + 1}{ext}")
        else:
            out = sibling(src, f"{unit} {i + 1}" if n > 1 else "", ext)
        save_image(im.convert("RGB"), out, target)
        outs.append(out)
    return outs


def pdf_to_images(src, target: str, dpi: int | None = None) -> list[Path]:
    dpi = dpi or int(CONFIG.get("pdf_dpi", 200))
    return _save_pages(_pdf_pages_to_images(src, dpi), src, target, "page")


def pdf_to_txt(src) -> list[Path]:
    with open_pdf(src) as doc:
        text = "\n\f\n".join(page.get_text("text") for page in doc)
    out = sibling(src, ext=".txt")
    out.write_text(text, encoding="utf-8")
    return [out]


def pdf_to_docx(src) -> list[Path]:
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Inches
    d = Document()
    with open_pdf(src) as doc:
        for pno, page in enumerate(doc):
            if pno:
                d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            items = []
            for b in page.get_text("dict")["blocks"]:
                y = b["bbox"][1]
                if b.get("type") == 0:
                    lines, size, bold = [], 11.0, False
                    for ln in b["lines"]:
                        spans = ln["spans"]
                        if spans:
                            bold = bold or any("Bold" in s["font"] for s in spans)
                        lines.append("".join(s["text"] for s in spans))
                    txt = " ".join(l.strip() for l in lines).strip()
                    if txt:
                        items.append((y, "text", (txt, size, bold)))
                elif b.get("type") == 1 and b.get("image"):
                    items.append((y, "img", (b["image"], b["bbox"])))
            for _, kind, data in sorted(items, key=lambda t: t[0]):
                if kind == "text":
                    txt, _size, bold = data
                    p = d.add_paragraph()
                    p.add_run(txt).bold = bold and len(txt) < 120
                else:
                    raw, bbox = data
                    try:
                        d.add_picture(io.BytesIO(raw), width=Inches(min(6.5, max(0.5, (bbox[2] - bbox[0]) / 72))))
                    except Exception:
                        pass
    out = sibling(src, ext=".docx")
    d.save(out)
    return [out]


def merge_pdfs(srcs: list) -> list[Path]:
    if len(srcs) < 2:
        raise UserError("Merge needs at least two PDFs.")
    out = sibling(srcs[0], "Merged", ".pdf")
    nd = fitz.open()
    for s in srcs:
        with open_pdf(s) as d:
            nd.insert_pdf(d)
    nd.save(out, garbage=3, deflate=True)
    nd.close()
    return [out]


def split_pdf(src, pages: list[int] | None = None) -> list[Path]:
    """Split a PDF into one file per page. `pages` (1-based) limits it to the chosen
    pages; None means every page. Multiple outputs go into a named folder."""
    src = Path(src)
    with open_pdf(src) as doc:
        n = doc.page_count
        idxs = [p - 1 for p in pages] if pages else list(range(n))
        idxs = [i for i in idxs if 0 <= i < n]
        if not idxs:
            raise UserError("No valid pages selected.")
        folder = sibling_folder(src) if len(idxs) > 1 else None
        outs = []
        for i in idxs:
            nd = fitz.open()
            nd.insert_pdf(doc, from_page=i, to_page=i)
            out = unique_path(folder / f"page {i + 1}.pdf") if folder else sibling(src, f"page {i + 1}", ".pdf")
            nd.save(out)
            nd.close()
            outs.append(out)
    return outs


def rotate_pdf(src, degrees: int = 90) -> list[Path]:
    out = sibling(src, "Rotated", ".pdf")
    with open_pdf(src) as doc:
        for page in doc:
            page.set_rotation((page.rotation + degrees) % 360)
        doc.save(out)
    return [out]


def compress_pdf(src) -> list[Path]:
    out = sibling(src, "Compressed", ".pdf")
    with open_pdf(src) as doc:
        try:
            doc.rewrite_images(dpi_threshold=160, dpi_target=150, quality=70,
                               lossy=True, lossless=True, color=True, gray=True)
        except Exception:
            pass  # older PyMuPDF: structural compression below still applies
        doc.save(out, garbage=4, deflate=True, deflate_images=True, clean=True)
    return [out]


# ------------------------------------------------------------------ docx / text
CSS = """* { font-family: sans-serif; } body { font-size: 11pt; line-height: 1.35; }
h1 { font-size: 20pt; } h2 { font-size: 16pt; } p { margin: 0 0 6pt 0; }
pre { font-family: monospace; font-size: 9.5pt; white-space: pre-wrap; }
table { border-collapse: collapse; } td { border: 1px solid #888; padding: 3pt 5pt; }"""


def _html_to_pdf(body_html: str, out: Path) -> None:
    story = fitz.Story(html=body_html, user_css=CSS)
    writer = fitz.DocumentWriter(str(out))
    mediabox = fitz.paper_rect("letter")
    where = mediabox + (54, 54, -54, -54)
    more = True
    while more:
        dev = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()


def _find_soffice() -> str | None:
    """LibreOffice headless binary, for high-fidelity office conversions when installed."""
    import shutil
    exe = shutil.which("soffice") or shutil.which("soffice.exe")
    if exe:
        return exe
    if sys.platform == "win32":
        for p in (r"C:\Program Files\LibreOffice\program\soffice.exe",
                  r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
            if os.path.exists(p):
                return p
    return None


def soffice_available() -> bool:
    return _find_soffice() is not None


def _soffice_to(src, target_ext: str) -> Path | None:
    """Convert any office document via LibreOffice headless. Safe: an argument array,
    never a shell string. Returns the output path, or None if LO is missing / it failed."""
    soffice = _find_soffice()
    if not soffice:
        return None
    import shutil
    import subprocess
    import tempfile
    src = Path(src)
    with tempfile.TemporaryDirectory() as tmp:
        try:
            subprocess.run(
                [soffice, "--headless", "--norestore", "--convert-to", target_ext,
                 "--outdir", tmp, str(src.resolve())],
                check=True, timeout=180, capture_output=True,
                creationflags=0x08000000 if sys.platform == "win32" else 0,  # no console flash
            )
        except Exception:
            return None
        produced = Path(tmp) / (src.stem + "." + target_ext)
        if not produced.exists():
            return None
        out = sibling(src, ext="." + target_ext)
        shutil.move(str(produced), str(out))
        return out


def _docx_to_html(src) -> str:
    from docx import Document
    d = Document(src)
    parts = []
    for p in d.paragraphs:
        body = html.escape(p.text) or "&nbsp;"
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style.startswith("heading"):
            lvl = min(int("".join(c for c in style if c.isdigit()) or "1"), 6)
            parts.append(f"<h{lvl}>{body}</h{lvl}>")
        elif style == "title":
            parts.append(f"<h1>{body}</h1>")
        else:
            parts.append(f"<p>{body}</p>")
    return "\n".join(parts)


def docx_to_pdf(src) -> list[Path]:
    src = Path(src)
    out = _soffice_to(src, "pdf")  # best fidelity, handles .doc/.odt/.rtf too
    if out:
        return [out]
    if src.suffix.lower() != ".docx":
        raise UserError(f"Converting {src.suffix} files needs LibreOffice installed.")
    out = sibling(src, ext=".pdf")  # pure-Python fallback for .docx
    _html_to_pdf(_docx_to_html(src), out)
    return [out]


def docx_to_txt(src) -> list[Path]:
    src = Path(src)
    if src.suffix.lower() == ".docx":
        from docx import Document
        d = Document(src)
        out = sibling(src, ext=".txt")
        out.write_text("\n".join(p.text for p in d.paragraphs), encoding="utf-8")
        return [out]
    out = _soffice_to(src, "txt")
    if out:
        return [out]
    raise UserError(f"Converting {src.suffix} files needs LibreOffice installed.")


def _read_text(src) -> str:
    raw = Path(src).read_bytes()
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def text_to_pdf(src) -> list[Path]:
    out = sibling(src, ext=".pdf")
    _html_to_pdf(f"<pre>{html.escape(_read_text(src))}</pre>", out)
    return [out]


def text_to_docx(src) -> list[Path]:
    from docx import Document
    d = Document()
    for line in _read_text(src).splitlines():
        d.add_paragraph(line)
    out = sibling(src, ext=".docx")
    d.save(out)
    return [out]


# ------------------------------------------------------------------ spreadsheets
def _read_sheet_rows(src) -> list[list]:
    """First worksheet as rows of cell values. Handles xlsx / xls / csv / tsv.
    (.ods is routed through LibreOffice instead.)"""
    ext = Path(src).suffix.lower()
    if ext in (".csv", ".tsv"):
        import csv
        delim = "\t" if ext == ".tsv" else ","
        with open(src, newline="", encoding="utf-8-sig", errors="replace") as f:
            return [row for row in csv.reader(f, delimiter=delim)]
    if ext == ".xls":
        import xlrd  # type: ignore
        sh = xlrd.open_workbook(src).sheet_by_index(0)
        return [[sh.cell_value(r, c) for c in range(sh.ncols)] for r in range(sh.nrows)]
    if ext == ".ods":
        raise UserError("Converting .ods files needs LibreOffice installed.")
    import openpyxl
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    rows = [["" if v is None else v for v in row] for row in wb.active.iter_rows(values_only=True)]
    wb.close()
    return rows


def _rows_to_html(rows: list[list]) -> str:
    def cell(v):
        return f"<td>{html.escape('' if v is None else str(v))}</td>"
    body = "".join("<tr>" + "".join(cell(v) for v in r) + "</tr>" for r in rows)
    return f"<table>{body}</table>"


def sheet_to_pdf(src) -> list[Path]:
    src = Path(src)
    out = _soffice_to(src, "pdf")  # keeps real formatting when LO is present
    if out:
        return [out]
    out = sibling(src, ext=".pdf")
    _html_to_pdf(_rows_to_html(_read_sheet_rows(src)), out)
    return [out]


def sheet_to_csv(src) -> list[Path]:
    import csv
    out = sibling(Path(src), ext=".csv")
    with open(out, "w", newline="", encoding="utf-8") as f:
        csv.writer(f).writerows(_read_sheet_rows(src))
    return [out]


def sheet_to_xlsx(src) -> list[Path]:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in _read_sheet_rows(src):
        ws.append(list(r))
    out = sibling(Path(src), ext=".xlsx")
    wb.save(out)
    return [out]


# ------------------------------------------------------------------ presentations
def slide_to_txt(src) -> list[Path]:
    src = Path(src)
    if src.suffix.lower() == ".pptx":
        from pptx import Presentation  # type: ignore
        lines = []
        for i, slide in enumerate(Presentation(src).slides, 1):
            lines.append(f"--- Slide {i} ---")
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    lines.append(shape.text_frame.text)
        out = sibling(src, ext=".txt")
        out.write_text("\n".join(lines), encoding="utf-8")
        return [out]
    out = _soffice_to(src, "txt")
    if out:
        return [out]
    raise UserError(f"Converting {src.suffix} to text needs LibreOffice installed.")


def _slide_images(src) -> list[Image.Image]:
    """Every slide as a Pillow image — via LibreOffice if present, else the built-in
    .pptx renderer."""
    src = Path(src)
    pdf = _soffice_to(src, "pdf")
    if pdf:
        try:
            return _pdf_pages_to_images(pdf, int(CONFIG.get("pdf_dpi", 150)))
        finally:
            try:
                Path(pdf).unlink()
            except OSError:
                pass
    if src.suffix.lower() != ".pptx":
        raise UserError(f"Converting {src.suffix} files needs LibreOffice installed.")
    from .render import pptx_to_images
    imgs = pptx_to_images(src)
    if not imgs:
        raise UserError("This presentation has no slides.")
    return imgs


def slide_to_pdf(src) -> list[Path]:
    src = Path(src)
    out = _soffice_to(src, "pdf")  # best fidelity when LibreOffice is present
    if out:
        return [out]
    imgs = [im.convert("RGB") for im in _slide_images(src)]
    out = sibling(src, ext=".pdf")
    imgs[0].save(out, "PDF", resolution=150.0, save_all=True, append_images=imgs[1:])
    return [out]


def slide_to_images(src, target: str = "PNG") -> list[Path]:
    return _save_pages(_slide_images(src), Path(src), target, "slide")


def split_slides(src, slides: list[int] | None = None) -> list[Path]:
    """Separate chosen slides, each as its own one-page PDF. `slides` is 1-based;
    None means every slide. Multiple outputs go into a named folder."""
    src = Path(src)
    imgs = _slide_images(src)
    n = len(imgs)
    idxs = [s - 1 for s in slides] if slides else list(range(n))
    idxs = [i for i in idxs if 0 <= i < n]
    if not idxs:
        raise UserError("No valid slides selected.")
    folder = sibling_folder(src) if len(idxs) > 1 else None
    outs = []
    for i in idxs:
        out = unique_path(folder / f"slide {i + 1}.pdf") if folder else sibling(src, f"slide {i + 1}", ".pdf")
        imgs[i].convert("RGB").save(out, "PDF", resolution=150.0)
        outs.append(out)
    return outs


def page_count(cat: str, src) -> int:
    """Number of pages (PDF) or slides (presentation), for the split page-picker."""
    if cat == "pdf":
        with open_pdf(src) as doc:
            return doc.page_count
    if cat == "slide":
        if Path(src).suffix.lower() == ".pptx":
            try:
                from pptx import Presentation
                return len(Presentation(str(src)).slides)
            except Exception:
                pass
        return len(_slide_images(src))
    return 1


# ------------------------------------------------------------------ action catalog
_TARGETS = {"image": IMAGE_TARGETS, "pdf": PDF_TARGETS, "docx": DOCX_TARGETS,
            "text": TEXT_TARGETS, "sheet": SHEET_TARGETS, "slide": SLIDE_TARGETS}


def convert_targets(cat: str, src_ext: str) -> list[str]:
    same = {"jpeg": "JPG", "jpg": "JPG", "tif": "TIFF", "tiff": "TIFF"}
    cur = same.get(src_ext.lower().lstrip("."), src_ext.upper().lstrip("."))
    return [t for t in _TARGETS.get(cat, []) if t != cur]


def _convert(srcs, cat, target):
    src = srcs[0]
    target = target.upper()
    if cat == "image":
        return image_to(src, target)
    if cat == "pdf":
        if target in ("JPG", "PNG"):
            return pdf_to_images(src, target)
        return pdf_to_txt(src) if target == "TXT" else pdf_to_docx(src)
    if cat == "docx":
        return docx_to_pdf(src) if target == "PDF" else docx_to_txt(src)
    if cat == "text":
        return text_to_pdf(src) if target == "PDF" else text_to_docx(src)
    if cat == "sheet":
        if target == "PDF":
            return sheet_to_pdf(src)
        return sheet_to_csv(src) if target == "CSV" else sheet_to_xlsx(src)
    if cat == "slide":
        if target == "PDF":
            return slide_to_pdf(src)
        if target in ("PNG", "JPG"):
            return slide_to_images(src, target)
        return slide_to_txt(src)
    raise UserError(f"Can't convert {cat} to {target}")


# Tool ops by category: (label, callable(srcs) -> list[Path], needs_multiple)
TOOLS = {
    "image": [
        ("Compress", lambda s: compress_image(s[0]), False),
        ("Resize", lambda s: resize_image(s[0]), False),
        ("→ PDF", images_to_pdf, True),
    ],
    "pdf": [
        ("Merge", merge_pdfs, True),
        ("Split", lambda s: split_pdf(s[0]), False),
        ("Rotate", lambda s: rotate_pdf(s[0]), False),
        ("Compress", lambda s: compress_pdf(s[0]), False),
    ],
    "docx": [],
    "text": [],
    "sheet": [],
    "slide": [
        ("Split", lambda s: split_slides(s[0]), False),
    ],
}

# Tools that ask the user to choose pages/slides first (handled in app.py).
NEEDS_PAGE_PICK = {"Split"}


def actions_for(cat: str, src_ext: str, n_files: int) -> dict[str, list[str]]:
    """Slice labels for the two rings, given the dropped file category and count."""
    convert = convert_targets(cat, src_ext)
    tools = [label for (label, _fn, multi) in TOOLS.get(cat, []) if (n_files > 1) == multi or not multi]
    if n_files > 1:
        # multi-select: only ops that take many files make sense
        convert = []
        tools = [label for (label, _fn, multi) in TOOLS.get(cat, []) if multi]
    return {"convert": convert, "tools": tools}


def run(cat: str, ring: str, label: str, srcs: list, params: dict | None = None) -> list[Path]:
    """Execute the chosen slice. `ring` is 'convert' or 'tools'. `params` carries
    options such as the selected pages for Split."""
    params = params or {}
    if ring == "convert":
        return _convert(srcs, cat, label)
    if label == "Split":
        pages = params.get("pages")  # 1-based list, or None for all
        if cat == "pdf":
            return split_pdf(srcs[0], pages)
        if cat == "slide":
            return split_slides(srcs[0], pages)
    for lbl, fn, _multi in TOOLS.get(cat, []):
        if lbl == label:
            return fn(srcs)
    raise UserError(f"Unknown action: {label}")
