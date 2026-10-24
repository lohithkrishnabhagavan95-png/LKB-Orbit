"""End-to-end checks of the processing engine on generated files."""
from pathlib import Path

import pytest
from PIL import Image

from quickconvert import engine


@pytest.fixture
def png(tmp_path) -> Path:
    p = tmp_path / "pic.png"
    Image.new("RGB", (120, 80), (200, 100, 50)).save(p)
    return p


def test_image_convert_compress_resize(png):
    (jpg,) = engine.image_to(png, "JPG")
    assert jpg.exists() and Image.open(jpg).format == "JPEG"
    (c,) = engine.compress_image(png, quality=50)
    assert c.exists()
    (r,) = engine.resize_image(png, long_side=40)
    assert max(Image.open(r).size) == 40
    assert png.exists()  # original preserved


def test_images_to_pdf_and_pdf_ops(png, tmp_path):
    png2 = tmp_path / "pic2.png"
    Image.new("RGB", (60, 60), (10, 20, 30)).save(png2)

    (pdf1,) = engine.images_to_pdf([png, png2])
    import pymupdf as fitz
    assert fitz.open(pdf1).page_count == 2

    (pdf2,) = engine.images_to_pdf([png2])
    (merged,) = engine.merge_pdfs([pdf1, pdf2])
    assert fitz.open(merged).page_count == 3

    pages = engine.split_pdf(pdf1)
    assert len(pages) == 2

    (rot,) = engine.rotate_pdf(pdf1, 90)
    assert fitz.open(rot)[0].rotation == 90

    imgs = engine.pdf_to_images(pdf1, "PNG")
    assert len(imgs) == 2 and all(p.exists() for p in imgs)

    (txt,) = engine.pdf_to_txt(pdf1)
    assert txt.exists()


def test_merge_requires_two(png):
    (pdf,) = engine.images_to_pdf([png])
    with pytest.raises(engine.UserError):
        engine.merge_pdfs([pdf])


def test_text_conversions(tmp_path):
    t = tmp_path / "note.txt"
    t.write_text("hello\nworld", encoding="utf-8")
    (pdf,) = engine.text_to_pdf(t)
    assert pdf.exists()
    (docx,) = engine.text_to_docx(t)
    assert docx.exists()


def test_docx_to_pdf(tmp_path):
    # Regression: docx->pdf runs Word via COM (or falls back to the built-in
    # renderer). It must produce a PDF without a hard COM crash.
    from docx import Document
    src = tmp_path / "doc.docx"
    d = Document()
    d.add_heading("Title", 0)
    d.add_paragraph("Some body text.")
    d.save(src)
    (pdf,) = engine.docx_to_pdf(src)
    assert pdf.exists() and pdf.suffix == ".pdf"


def test_sheet_conversions(tmp_path):
    import openpyxl
    xlsx = tmp_path / "data.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["name", "qty"])
    ws.append(["apples", 5])
    wb.save(xlsx)

    (csv,) = engine.sheet_to_csv(xlsx)
    assert csv.exists() and "apples" in csv.read_text(encoding="utf-8")
    (pdf,) = engine.sheet_to_pdf(xlsx)
    assert pdf.exists() and pdf.suffix == ".pdf"
    (x2,) = engine.sheet_to_xlsx(csv)   # csv -> xlsx round-trips
    assert x2.exists() and x2.suffix == ".xlsx"
    assert set(engine.convert_targets("sheet", ".xlsx")) == {"PDF", "CSV"}
    assert "XLSX" in engine.convert_targets("sheet", ".csv")


def test_slide_to_txt(tmp_path):
    from pptx import Presentation
    p = tmp_path / "deck.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    slide.shapes.title.text = "Hello Deck"
    prs.save(p)
    (txt,) = engine.slide_to_txt(p)
    assert txt.exists() and "Hello Deck" in txt.read_text(encoding="utf-8")


def test_pptx_to_pdf_and_images(tmp_path):
    # Self-contained render (no LibreOffice): multi-slide pptx -> PDF and PNGs.
    from pptx import Presentation
    import pymupdf as fitz
    p = tmp_path / "deck.pptx"
    prs = Presentation()
    for t in ("One", "Two", "Three"):
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = t
    prs.save(p)

    (pdf,) = engine.slide_to_pdf(p)
    assert pdf.exists() and fitz.open(pdf).page_count == 3   # one page per slide

    pngs = engine.slide_to_images(p, "PNG")
    assert len(pngs) == 3 and all(x.exists() for x in pngs)


def test_multipage_images_go_in_a_folder(tmp_path):
    from PIL import Image
    pngs = []
    for i in range(3):
        p = tmp_path / f"{i}.png"
        Image.new("RGB", (40, 30), (10 * i, 20, 30)).save(p)
        pngs.append(p)
    (pdf,) = engine.images_to_pdf(pngs)
    outs = engine.pdf_to_images(pdf, "PNG")
    assert len(outs) == 3
    parents = {p.parent for p in outs}
    assert len(parents) == 1                      # all pages in one folder
    folder = parents.pop()
    assert folder.is_dir() and folder != pdf.parent   # a dedicated named subfolder


def test_split_selected_pages(tmp_path):
    from PIL import Image
    pngs = []
    for i in range(4):
        p = tmp_path / f"{i}.png"
        Image.new("RGB", (30, 30), (i * 40, 0, 0)).save(p)
        pngs.append(p)
    (pdf,) = engine.images_to_pdf(pngs)
    outs = engine.split_pdf(pdf, [1, 3])          # only pages 1 and 3
    assert len(outs) == 2 and all(o.exists() for o in outs)


def test_split_slides_selected(tmp_path):
    from pptx import Presentation
    p = tmp_path / "d.pptx"
    pr = Presentation()
    for t in ("A", "B", "C"):
        s = pr.slides.add_slide(pr.slide_layouts[5])
        s.shapes.title.text = t
    pr.save(p)
    assert engine.page_count("slide", p) == 3
    outs = engine.split_slides(p, [2])            # just slide 2
    assert len(outs) == 1 and outs[0].exists() and outs[0].suffix == ".pdf"


def test_heic_avif_targets():
    if not engine.HEIF_OK:
        import pytest
        pytest.skip("pillow-heif not installed")
    from PIL import Image
    import tempfile
    from pathlib import Path
    d = Path(tempfile.mkdtemp())
    png = d / "p.png"
    Image.new("RGB", (40, 30), (10, 20, 30)).save(png)
    (heic,) = engine.image_to(png, "HEIC")
    assert heic.exists() and heic.suffix == ".heic"
    # and back
    (jpg,) = engine.image_to(heic, "JPG")
    assert jpg.exists()
    assert "HEIC" in engine.convert_targets("image", ".png")


def test_epub_conversions(tmp_path):
    from ebooklib import epub
    book = epub.EpubBook()
    book.set_title("Test Book")
    c = epub.EpubHtml(title="Ch1", file_name="c1.xhtml")
    c.content = "<html><body><h1>Chapter One</h1><p>Hello world from an ebook.</p></body></html>"
    book.add_item(c)
    book.spine = [c]
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    p = tmp_path / "book.epub"
    epub.write_epub(str(p), book)

    (txt,) = engine.epub_to_txt(p)
    assert txt.exists() and "Hello world" in txt.read_text(encoding="utf-8")
    (pdf,) = engine.epub_to_pdf(p)
    assert pdf.exists() and pdf.suffix == ".pdf"
    from quickconvert.util import category
    assert category("book.epub") == "ebook"


def test_categories():
    from quickconvert.util import category
    assert category("a.xlsx") == "sheet"
    assert category("a.csv") == "sheet"
    assert category("a.pptx") == "slide"
    assert category("a.docx") == "docx"
    assert category("a.png") == "image"
    assert category("a.heic") == "image"
    assert category("a.epub") == "ebook"


def test_actions_catalog():
    a = engine.actions_for("image", ".png", 1)
    assert "JPG" in a["convert"] and "Compress" in a["tools"]
    assert engine.actions_for("pdf", ".pdf", 2)["tools"] == ["Merge"]
